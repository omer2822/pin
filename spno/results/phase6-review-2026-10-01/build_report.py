from pathlib import Path
import re, json, statistics
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.70);sec.right_margin=Inches(.70)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.30)
for name,size in [('Normal',11),('Title',22),('Subtitle',12),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
 s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 for border in s.element.xpath('.//w:pBdr'): border.getparent().remove(border)
 szcs=s.element.get_or_add_rPr().find(qn('w:szCs'))
 if szcs is None: szcs=OxmlElement('w:szCs');s.element.get_or_add_rPr().append(szcs)
 szcs.set(qn('w:val'),str(int(size*2)))
 s.paragraph_format.space_after=Pt(6)
 s.paragraph_format.line_spacing=1.12
 if name.startswith('Heading'):
  s.paragraph_format.space_before=Pt(12);s.paragraph_format.keep_with_next=True
 s.element.get_or_add_rPr().append(OxmlElement('w:rtl'))
 s.element.get_or_add_rPr().find(qn('w:rtl')).set(qn('w:val'),'0')
 doc.styles[name].element.get_or_add_rPr().find(qn('w:rFonts')).set(qn('w:cs'),'Arial')

def direction(p,rtl=True):
 e=OxmlElement('w:bidi');e.set(qn('w:val'),'1' if rtl else '0');p._p.get_or_add_pPr().append(e)
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT

def clean(s):
 return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',s)

def text_runs(p,s,size=None,color=None):
 # Keep Word's bidirectional Unicode algorithm; specify Hebrew and Latin run direction.
 for segment in re.split(r'(\[[^\]]+\]\([^)]+\))',s):
  link=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',segment)
  if link:
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(link.group(2),RT.HYPERLINK,is_external=True))
   r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');rp.append(c)
   u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp)
   t=OxmlElement('w:t');t.text=link.group(1);r.append(t);h.append(r);p._p.append(h)
  elif segment: plain_runs(p,segment,size,color)

def plain_runs(p,s,size=None,color=None):
 for part in re.split(r'(\*\*.*?\*\*)',s):
  bold=part.startswith('**') and part.endswith('**');part=part[2:-2] if bold else part
  for tok in re.split(r'([\u0590-\u05ff][\u0590-\u05ff\s־״׳]*)',part):
   if not tok:continue
   r=p.add_run(tok);r.bold=bold;r.font.name='Arial'
   if size:r.font.size=Pt(size)
   if color:r.font.color.rgb=RGBColor.from_string(color)
   rp=r._r.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
   if rf is not None:rf.set(qn('w:cs'),'Arial')
   e=OxmlElement('w:rtl');e.set(qn('w:val'),'1' if re.search('[\u0590-\u05ff]',tok) else '0');rp.append(e)

def paragraph(s,style=None):
 p=doc.add_paragraph(style=style);direction(p,bool(re.search('[\u0590-\u05ff]',s)));text_runs(p,s);return p

def table(lines):
 rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
 rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in r)]
 n=len(rows[0]);assert all(len(r)==n for r in rows), 'Malformed table row'
 t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths={2:[1.7,5.17],3:[1.25,2.81,2.81],4:[1.25,1.9,1.9,1.82],5:[1.5,1.34,1.34,1.34,1.35],6:[1.20,1.13,1.13,1.13,1.14,1.14]}[n]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 pr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for edge in ['top','bottom','left','right']:
  e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
 pr.append(margins)
 if re.search('[\u0590-\u05ff]',rows[0][0]):
  e=OxmlElement('w:bidiVisual');pr.append(e)
 for idx,row in enumerate(rows):
  cells=t.add_row().cells
  for j,(cell,s) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];direction(p,bool(re.search('[\u0590-\u05ff]',s)))
   if not re.search('[\u0590-\u05ff]',s):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
   text_runs(p,s,size=9.2 if n>=5 else 9.6,color='FFFFFF' if idx==0 else '000000')
   if idx==0:
    for r in p.runs:r.bold=True
   tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'283E50' if idx==0 else ('F3F6F8' if idx%2==0 else 'FFFFFF'));tcpr.append(shade)
  if idx==0:
   trpr=t.rows[-1]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  cant=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(cant)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)

lines=(ROOT/'research-review-he.md').read_text().splitlines();i=0
while i<len(lines):
 l=lines[i]
 if l.startswith('|'):
  batch=[]
  while i<len(lines) and lines[i].startswith('|'):batch.append(lines[i]);i+=1
  table(batch);continue
 if not l.strip():i+=1;continue
 if l.startswith('# '):paragraph(l[2:],'Title')
 elif l.startswith('## '):paragraph(l[3:],'Heading 1')
 elif l.startswith('### '):paragraph(l[4:],'Heading 2')
 else:paragraph(l)
 i+=1

header=sec.header.paragraphs[0];direction(header);text_runs(header,'PINO SPNO   ביקורת Phase 6   1 באוקטובר 2026',size=8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
text_runs(footer,'Phase 6 research review  |  ',size=8)
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='ביקורת מחקרית של תוצאות Phase 6 בפרויקט PINO SPNO'
doc.core_properties.subject='Experimental fairness spectral identifiability and C1 failure modes'
doc.core_properties.author='Scientific research review'
doc.save(ROOT/'Phase6_research_review_HE.docx')
print('Created',ROOT/'Phase6_research_review_HE.docx')
